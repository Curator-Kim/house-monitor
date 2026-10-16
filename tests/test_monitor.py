import json
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from house_monitor.models import Observation, SourceResult
from house_monitor.storage import Store
from house_monitor.schedule import daily_target, scheduler_enabled, tick
from house_monitor.parsers import price, fang_list, fang_history, lianjia_list, official_json, ygjy_records, nbs_indices
from house_monitor.http import robots_allowed, blocked, SourceError, Client
from house_monitor.analytics import paired_changes, csv_bytes, excel_bytes
from house_monitor.sources import collect_list
from house_monitor.replay import replay

TZ = ZoneInfo("Asia/Shanghai")
REGION = {"id": "440112", "city": "广州", "district": "黄埔"}
URL = "https://gz.newhouse.fang.com/house/s/huangpu/"


@pytest.fixture
def config(tmp_path):
    return {"data_dir": str(tmp_path), "timezone": "Asia/Shanghai", "schedule_start": "09:00:00",
            "schedule_end": "12:00:00", "max_pages": 30, "history_enabled": False,
            "request_interval_seconds": 0, "request_timeout_seconds": 1, "regions": [REGION]}


def card(pid="1", name="测试楼盘", tags="普通住宅", status="在售", area="黄埔", raw="22000元/㎡(单价)"):
    return f'''<li id="lp_{pid}"><div class="nlcd_name"><a href="/loupan/{pid}.htm">{name}</a></div>
    <div class="address"><a title="测试地址">[{area}] 测试地址</a></div><div class="fangyuan">
    <span class="{'inSale' if status == '在售' else 'soldOut'}">{status}</span><a>{tags}</a></div>
    <div class="nhouse_price">{raw}</div></li>'''


def page(cards, number=1, total=1, nxt=None):
    return f'<div id="newhouse_loupan_list"><ul>{cards}</ul></div><div>{number}/{total}</div>' + (f'<a href="{nxt}">下一页</a>' if nxt else "")


@pytest.mark.parametrize("raw,value,high,kind,unit", [
    ("22,000元/㎡", 22000, None, "reference_unit", "元/平方米"),
    ("2.1万元/平方米起", 21000, None, "reference_unit_starting", "元/平方米"),
    ("20000-24000元/平方米", 20000,24000,"reference_unit_range","元/平方米"),
    ("280万元/套起(总价)",2800000,None,"reference_total_starting","元/套"),
    ("价格待定",None,None,"reference_unit","元/平方米"),
])
def test_price(raw,value,high,kind,unit):
    result = price(raw)
    assert (result["value"],result["value_high"],result["kind"],result["unit"]) == (value,high,kind,unit)


def test_filter_dedup_and_unknown():
    html = page(card() + card() + card("2",tags="写字楼") + card("3",status="售完") + card("4",area="番禺") + card("5",raw="价格待定"))
    rows, discovered, nxt, total = fang_list(html,URL,REGION)
    assert len(discovered) == 5 and len(rows) == 2
    assert rows[1].value is None and rows[1].missing_reason
    assert price("280万")["value"] is None


def test_history_separate_dates():
    html = '<ul><li><span>2025-11-20</span><span>22000元/平方米</span>带装修</li><li><span>2024-11-17</span><span>36000元/平方米</span></li></ul>'
    rows = fang_history(html,"1","楼盘",URL,"440112")
    assert len(rows) == 2
    assert all(r.granularity == "event" and r.period == r.source_updated for r in rows)


def test_nbs_month_city_normalization():
    html = '<h1>2026年8月份70个大中城市商品住宅销售价格变动情况</h1><p>表1：新建</p><table><tr><td>广 州</td><td>100.1</td><td>98.1</td><td>97.4</td></tr></table><p>表2：二手</p><table><tr><td>广州</td><td>100</td><td>96.2</td><td>95</td></tr></table>'
    rows = nbs_indices(html,URL,REGION)
    assert len(rows)==6 and {r.period for r in rows}=={"2026-08"}
    assert {r.kind for r in rows} >= {"new_index_mom","resale_index_mom"}


def test_robot_wildcard_and_longest_allow():
    rules = "User-agent: *\nDisallow: /house/$\nDisallow: /house/s/*-*-*-*-*\nDisallow: /private\nAllow: /private/public"
    assert robots_allowed(rules,URL)
    assert not robots_allowed(rules,"https://x/house/")
    assert not robots_allowed(rules,"https://x/house/s/a-b-c-d-e/")
    assert robots_allowed(rules,"https://x/private/public")
    assert not robots_allowed(rules,"https://x/private/secret")


def test_persistent_schedule_bounds_and_next_day(config):
    store = Store(config["data_dir"])
    current = datetime(2026,10,12,8,tzinfo=TZ)
    a = daily_target(store,config,current,lambda n: n-1)
    assert datetime.fromisoformat(a["target"]).hour == 12
    b = daily_target(store,config,current,lambda n: 0)
    assert a == b
    c = daily_target(store,config,current+timedelta(days=1),lambda n: 0)
    assert datetime.fromisoformat(c["target"]).hour == 9 and c["day"] != a["day"]


def test_tick_wait_catchup_once_and_no_backfill(config):
    store = Store(config["data_dir"])
    morning = datetime(2026,10,12,8,tzinfo=TZ)
    assert tick(config,morning)["state"] == "waiting"
    calls=[]
    def collect(*args,**kwargs):
        calls.append(kwargs)
        return {"run_id":99,"status":"partial"}
    assert tick(config,morning.replace(hour=15),collect)["state"] == "done"
    assert tick(config,morning.replace(hour=16),collect)["state"] == "done"
    assert len(calls)==1
    tick(config,morning+timedelta(days=4,hours=7),collect)
    assert len(calls)==2 and len(store.rows("SELECT * FROM schedules"))==2


def test_interrupted_run_bounded_retry(config):
    current = datetime(2026,10,12,15,tzinfo=TZ)
    def fail(*args,**kwargs):
        raise RuntimeError("crash")
    for _ in range(2):
        with pytest.raises(RuntimeError):
            tick(config,current,fail)
    assert tick(config,current,fail)["state"]=="failed"


def test_store_dedup_missing_never_overwrites(config):
    store = Store(config["data_dir"])
    stamp="2026-10-12T10:00:00+08:00"
    o=Observation("fang","1","楼盘",URL,"440112",value=22000)
    result=SourceResult("fang","440112",records=[o],complete=True)
    store.save_result(1,result,stamp)
    store.save_result(2,result,stamp)
    assert len(store.observations())==1
    o.value=None
    store.save_result(3,result,stamp)
    assert store.observations()[0]["value"]==22000
    o.source="lianjia"
    o.value=24000
    store.save_result(4,SourceResult("lianjia","440112",records=[o]),stamp)
    assert len(store.observations())==2
    assert len(store.rows("SELECT * FROM projects"))==2


def test_exports_and_same_sample():
    df=pd.DataFrame([
        {"source":"fang","region_id":"440112","source_id":"1","name":"=bad","granularity":"snapshot","kind":"reference_unit","period":"2026-10-12","value":20000},
        {"source":"fang","region_id":"440112","source_id":"1","name":"=bad","granularity":"snapshot","kind":"reference_unit","period":"2026-10-13","value":22000},
        {"source":"fang","region_id":"440112","source_id":"2","name":"other","granularity":"snapshot","kind":"reference_unit","period":"2026-10-13","value":50000},
    ])
    result=paired_changes(df)
    assert result.iloc[0].sample_count==1 and result.iloc[0].change_pct==pytest.approx(10)
    wb=load_workbook(BytesIO(excel_bytes(df)))
    assert wb.active.max_row==4 and wb.active['D2'].value=="'=bad"
    assert "'=bad" in csv_bytes(df).decode("utf-8-sig")


class FakeClient:
    def __init__(self, pages):
        self.pages=pages
    def get(self,url):
        value=self.pages[url]
        if isinstance(value,Exception):
            raise value
        return value


def test_full_pagination_and_partial_failure(config):
    nxt=URL+'b92/'
    client=FakeClient({URL:page(card(),1,2,nxt),nxt:page(card('2'),2,2)})
    result=collect_list("fang",URL,REGION,client,config,Store(config["data_dir"]))
    assert result.complete and result.pages==2 and result.discovered==2
    client.pages[nxt]=SourceError("captcha","blocked")
    result=collect_list("fang",URL,REGION,client,config,Store(config["data_dir"]))
    assert not result.complete and result.status=="partial" and len(result.records)==1


def test_bad_pagination_not_silent_success(config):
    nxt=URL+'b92/'
    client=FakeClient({URL:page(card(),1,2,nxt),nxt:page(card(),2,2)})
    result=collect_list("fang",URL,REGION,client,config,Store(config["data_dir"]))
    assert not result.complete


def test_validation_requires_calendar_days_and_checks(config):
    store=Store(config["data_dir"])
    o=Observation("fang","1","楼盘",URL,"440112",value=22000)
    result=SourceResult("fang","440112",records=[o],complete=True)
    for i in range(3):
        store.save_result(i,result,f"2026-10-{12+i}T10:00:00+08:00")
    assert store.gates()[0]["stage"]=="provisional"
    with store.connect() as c:
        c.execute("INSERT INTO quality_checks VALUES ('fang','440112','1','2026-10-14','核对通过')")
    assert store.gates()[0]["stage"]=="verified"


def test_http_no_retry_after_access_restriction(config):
    import requests
    response=requests.Response()
    response.status_code=403
    response.url=URL
    response._content=b'Forbidden'
    client=Client(config,1)
    class Session:
        calls=0
        def get(self,*args,**kwargs):
            self.calls+=1
            return response
    client.session=Session()
    with pytest.raises(SourceError):
        client._request(URL)
    assert client.session.calls==1
    assert list(client.root.glob('*.body'))


def test_official_json_accepts_slash_date():
    payload={"data":[{"sectionName":"黄埔区","createTime":"2026/10/11 0:00:00",
                      "zhuZaiTaoShu":52,"zhuZaiArea":5600.2723}]}
    rows=official_json(payload,"signed",URL,REGION)
    assert len(rows)==2
    assert {r.kind for r in rows}=={"signed_count","signed_area"}
    assert all(r.period=="2026-10-11" and r.granularity=="day" for r in rows)


def test_replay_supports_lianjia_and_official(config):
    root=Path(config["data_dir"])/"raw"/"99"
    root.mkdir(parents=True)
    api="https://zfcj.gz.gov.cn/ysqgk/Api/WebApi/mrxjspfqyxx.ashx"
    payload={"data":[{"sectionName":"黄埔区","createTime":"2026/10/11 0:00:00",
                      "zhuZaiTaoShu":52,"zhuZaiArea":5600.2723}]}
    official_meta={"requested_url":api,"method":"POST","final_url":api,"status":200,
                   "fetched_at":"2026-10-12T04:51:18+00:00"}
    (root/"0001-official.json").write_text(json.dumps(official_meta),encoding="utf-8")
    (root/"0001-official.body").write_text(json.dumps(payload,ensure_ascii=False),encoding="utf-8")
    lianjia_url="https://gz.fang.lianjia.com/loupan/"
    html='''<div class="resblock-list"><div class="resblock-name"><a href="/loupan/p_test/">黄埔测试盘</a></div><div class="resblock-location">黄埔 / 科学城 / 测试路</div><div class="resblock-type">住宅</div><div class="sale-status">在售</div><div class="main-price">25000 元/㎡(均价)</div></div>'''
    lianjia_meta={"requested_url":lianjia_url,"method":"GET","final_url":lianjia_url,"status":200,
                  "fetched_at":"2026-10-12T04:51:07+00:00"}
    (root/"0002-lianjia.json").write_text(json.dumps(lianjia_meta),encoding="utf-8")
    (root/"0002-lianjia.body").write_text(html,encoding="utf-8")
    result=replay(config,99,"440112")
    assert {r["source"] for r in result["sources"]}=={"official","lianjia"}
    assert all(r["status"]=="ok" for r in result["sources"])
    rows=Store(config["data_dir"]).observations()
    assert {r["source"] for r in rows}=={"official","lianjia"}
    assert len(rows)==3

def test_replay_preserves_parse_failure_status(config):
    root=Path(config["data_dir"])/"raw"/"100"
    root.mkdir(parents=True)
    url="https://gz.fang.lianjia.com/loupan/"
    meta={"requested_url":url,"method":"GET","final_url":url,"status":200,
          "fetched_at":"2026-10-12T04:51:07+00:00"}
    (root/"0001-lianjia.json").write_text(json.dumps(meta),encoding="utf-8")
    (root/"0001-lianjia.body").write_text("<html></html>",encoding="utf-8")
    result=replay(config,100,"440112")
    assert len(result["sources"])==1
    assert result["sources"][0]["source"]=="lianjia"
    assert result["sources"][0]["status"]=="schema_changed"

def test_ygjy_project_records_keep_official_identity():
    payload={"currentPage":1,"totalPage":95,"totalNum":947,"data":[{
        "projectId":"official-project-id","projectName":"测试项目住宅","developer":"测试开发商",
        "developerId":"developer-id","presell":"20260066","projectAddress":"黄埔区测试路1号",
        "houseSoldNum":"6","houseUnsaleNum":"234"}]}
    rows=ygjy_records(payload,"https://example.test/ygjy",REGION,"2026-10-12")
    assert len(rows)==2
    assert {r.kind for r in rows}=={"sold_count","unsold_count"}
    assert {r.value for r in rows}=={6.0,234.0}
    assert all(r.source=="ygjy" and r.source_id=="official-project-id" and r.presell=="20260066" for r in rows)

def test_replay_supports_ygjy_pages(config):
    root=Path(config["data_dir"])/"raw"/"101"
    root.mkdir(parents=True)
    api="https://zfcj.gz.gov.cn/ysqgk/Api/WebApi/fdcxmxxlb.ashx"
    for page,project_id in ((1,"ygjy-1"),(2,"ygjy-2")):
        payload={"currentPage":page,"totalPage":2,"totalNum":2,"data":[{
            "projectId":project_id,"projectName":f"项目{page}","developer":"开发商","developerId":"dev",
            "presell":f"2026000{page}","projectAddress":"黄埔区测试路","houseSoldNum":"1","houseUnsaleNum":"2"}]}
        meta={"requested_url":api,"method":"POST","final_url":api,"status":200,
              "fetched_at":f"2026-10-12T05:0{page}:00+00:00"}
        (root/f"000{page}-ygjy.json").write_text(json.dumps(meta),encoding="utf-8")
        (root/f"000{page}-ygjy.body").write_text(json.dumps(payload,ensure_ascii=False),encoding="utf-8")
    result=replay(config,101,"440112")
    assert len(result["sources"])==1
    assert result["sources"][0]["status"]=="ok"
    assert result["sources"][0]["discovered"]==2
    rows=Store(config["data_dir"]).observations()
    assert len(rows)==4 and {r["source_id"] for r in rows}=={"ygjy-1","ygjy-2"}

def test_fang_citywide_filter_accepts_any_district():
    html = page(card("9", area="番禺"))
    rows, ids, nxt, total = fang_list(html, URL, {"id": "sz-all", "city": "深圳", "district": "全市"})
    assert len(rows) == 1 and rows[0].source_id == "9"

def test_watchlist_add_remove(config):
    store = Store(config["data_dir"])
    store.add_watch("fang", "440112", "1", "测试楼盘", "2026-10-13T10:00:00+08:00")
    rows = store.watchlist()
    assert len(rows) == 1 and rows[0]["name"] == "测试楼盘"
    store.add_watch("fang", "440112", "1", "更新名称", "2026-10-13T11:00:00+08:00")
    assert store.watchlist()[0]["name"] == "更新名称"
    store.remove_watch("fang", "440112", "1")
    assert store.watchlist() == []

def test_scheduler_node_switch(monkeypatch):
    monkeypatch.delenv("HOUSE_SCHEDULER_ENABLED", raising=False)
    assert scheduler_enabled({"scheduler_enabled": True}) is True
    assert scheduler_enabled({"scheduler_enabled": False}) is False
    monkeypatch.setenv("HOUSE_SCHEDULER_ENABLED", "1")
    assert scheduler_enabled({"scheduler_enabled": False}) is True
    monkeypatch.setenv("HOUSE_SCHEDULER_ENABLED", "0")
    assert scheduler_enabled({"scheduler_enabled": True}) is False